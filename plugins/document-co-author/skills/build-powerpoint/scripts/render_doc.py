#!/usr/bin/env python3
"""
render_doc.py — Word document renderer for .deck_spec.yaml content

Usage:
    python3 render_doc.py <spec.yaml>
    python3 render_doc.py <spec.yaml> --output <path.docx>
    python3 render_doc.py --check-repair <file.docx>

The spec's `doc_output` key controls the default output path.
Falls back to swapping .pptx → .docx in the `output` key.
"""

import sys
import yaml
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Brand colours ──────────────────────────────────────────────────────────
SN_TEAL   = RGBColor(0x00, 0xC7, 0xB1)
SN_ORANGE = RGBColor(0xF2, 0x6B, 0x43)
SN_SLATE  = RGBColor(0x5B, 0x7B, 0x8C)
SN_NAVY   = RGBColor(0x29, 0x3E, 0x40)
SN_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
SN_BLACK  = RGBColor(0x1A, 0x1A, 0x1A)
SN_GRAY   = RGBColor(0x55, 0x55, 0x55)
SN_LIGHT  = RGBColor(0xDD, 0xDD, 0xDD)

GROUP_COLOR = {
    "sn_teal":   SN_TEAL,
    "sn_orange": SN_ORANGE,
    "sn_slate":  SN_SLATE,
    "sn_navy":   SN_NAVY,
}
GROUP_HEX = {
    "sn_teal":   "00C7B1",
    "sn_orange": "F26B43",
    "sn_slate":  "5B7B8C",
    "sn_navy":   "293E40",
}
SCORE_HEX = {0: "8B0000", 1: "F26B43", 2: "0070C0", 3: "00C7B1"}


# ── XML helpers ─────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_bottom_border(para, hex_color="CCCCCC", size=4):
    """Draw a bottom rule under a paragraph."""
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), str(size))
    bot.set(qn("w:space"), "4")
    bot.set(qn("w:color"), hex_color)
    pBdr.append(bot)
    pPr.append(pBdr)


def clear_cell(cell):
    tc = cell._tc
    paras = tc.findall(qn("w:p"))
    for p in paras[1:]:
        tc.remove(p)
    for r in paras[0].findall(qn("w:r")):
        paras[0].remove(r)


def fill_cell(cell, text, bold=False, italic=False,
              color=None, size=None, align=None, shading=None):
    clear_cell(cell)
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    if shading:
        set_cell_bg(cell, shading)
    run = p.add_run(str(text))
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return p


def set_col_widths(table, widths_inches):
    for i, w in enumerate(widths_inches):
        table.columns[i].width = Inches(w)


# ── Paragraph helpers ───────────────────────────────────────────────────────
def para(doc, text="", size=11, color=None, bold=False, italic=False,
         before=0, after=6, indent=None, align=None):
    """Add a normal text paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    if align:
        p.alignment = align
    if text:
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        if color:
            run.font.color.rgb = color
    return p


def heading1(doc, text, color=None, hex_border=None, before=16, after=4):
    """Section heading — large, colored, with optional bottom rule."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if hex_border:
        add_bottom_border(p, hex_border, size=6)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = color or SN_NAVY
    return p


def heading2(doc, text, color=None, before=12, after=2):
    """Criterion heading — medium bold."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = color or SN_BLACK
    return p


def callout(doc, text, before=4, after=8):
    """Italic anecdote paragraph, indented."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.left_indent = Inches(0.28)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = SN_GRAY
    return p


def rule(doc):
    """Thin horizontal separator."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_bottom_border(p, "DDDDDD", size=4)
    return p


# ── Slide renderers ─────────────────────────────────────────────────────────
def render_cover(doc, slide):
    para(doc, before=0, after=60)  # top spacer

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(slide.get("title", "").replace("\n", " "))
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = SN_TEAL

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(slide.get("subtitle", ""))
    run.font.size = Pt(13)
    run.font.color.rgb = SN_NAVY

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(slide.get("date", ""))
    run.font.size = Pt(10)
    run.font.color.rgb = SN_SLATE

    doc.add_page_break()


def render_dimensions_overview(doc, slide):
    heading1(doc, slide.get("title", "Evaluation Dimensions"),
             color=SN_NAVY, hex_border="00C7B1", before=0)
    para(doc, slide.get("subtitle", ""), size=10, color=SN_GRAY, after=12)

    for section in slide.get("sections", []):
        color_key = section.get("color", "sn_teal")
        group_color = GROUP_COLOR.get(color_key, SN_TEAL)

        # Dimension label
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(section["label"])
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = group_color

        # Criteria as a bullet list
        for item in section.get("items", []):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(item)
            run.font.size = Pt(9)
            run.font.color.rgb = SN_BLACK

    doc.add_page_break()


def render_scoring_table(doc, slide):
    """Each criterion becomes a named section with prose, not a table row."""
    systems = slide.get("systems", [])

    for group in slide.get("groups", []):
        color_key = group.get("color", "sn_teal")
        group_color = GROUP_COLOR.get(color_key, SN_TEAL)
        hex_c = GROUP_HEX.get(color_key, "00C7B1")

        # ── Group heading ───────────────────────────────────────────────
        heading1(doc, group.get("name", ""),
                 color=group_color, hex_border=hex_c, before=20)

        for ci, cr in enumerate(group.get("criteria", [])):
            # ── Criterion heading ────────────────────────────────────────
            cr_id = cr.get("id", "")
            cr_name = cr.get("name", "")
            heading2(doc, f"{cr_id}  ·  {cr_name}", color=SN_NAVY)

            # Assessment question
            para(doc, cr.get("assess", ""), size=10, before=0, after=4)

            # Anecdote
            anecdote = cr.get("anecdote", "")
            if anecdote:
                callout(doc, anecdote)

            # Score input row when systems are specified
            if systems:
                score_row(doc, systems)

            # Separator between criteria (not after the last one)
            if ci < len(group.get("criteria", [])) - 1:
                rule(doc)


def score_row(doc, systems):
    """Inline score fields — only rendered when systems are named."""
    table = doc.add_table(rows=1, cols=len(systems))
    table.style = "Table Grid"
    for i, name in enumerate(systems):
        cell = table.rows[0].cells[i]
        clear_cell(cell)
        set_cell_bg(cell, "F5F5F5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{name}:  ___ / 3")
        run.font.size = Pt(9)
        run.font.color.rgb = SN_GRAY
        cell.width = Inches(7.0 / len(systems))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def render_scoring_guide(doc, slide):
    heading1(doc, slide.get("title", "How to Score"),
             color=SN_NAVY, hex_border="00C7B1", before=0)
    para(doc, slide.get("subtitle", ""), size=10, color=SN_GRAY, after=14)

    # ── Score definitions ────────────────────────────────────────────────
    rows_data = slide.get("rows", [])
    if rows_data:
        p = doc.add_paragraph()
        r = p.add_run("Score Definitions")
        r.bold = True
        r.font.size = Pt(12)
        p.paragraph_format.space_after = Pt(6)

        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        hdr = table.rows[0]
        for cell, label in zip(hdr.cells, ["Score", "Meaning", "What it means"]):
            fill_cell(cell, label, bold=True, color=SN_WHITE, size=10, shading="293E40")

        for i, row_data in enumerate(rows_data):
            row = table.add_row()
            bg = "F5F5F5" if i % 2 == 0 else "FFFFFF"
            score_str = str(row_data.get("score", ""))
            hex_c = SCORE_HEX.get(int(score_str) if score_str.isdigit() else -1, "000000")
            sc = RGBColor(int(hex_c[:2], 16), int(hex_c[2:4], 16), int(hex_c[4:], 16))
            fill_cell(row.cells[0], score_str, bold=True, color=sc, size=13,
                      align=WD_ALIGN_PARAGRAPH.CENTER, shading=bg)
            fill_cell(row.cells[1], row_data.get("meaning", ""), bold=True, size=10, shading=bg)
            fill_cell(row.cells[2], row_data.get("detail", ""), size=10, shading=bg)

        set_col_widths(table, [0.65, 1.6, 4.75])
        para(doc, after=16)

    # ── Dimension summary ────────────────────────────────────────────────
    weights = slide.get("weights", [])
    if weights:
        p = doc.add_paragraph()
        r = p.add_run("Scoring Summary")
        r.bold = True
        r.font.size = Pt(12)
        p.paragraph_format.space_after = Pt(6)

        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        hdr = table.rows[0]
        for cell, label in zip(hdr.cells, ["Dimension", "Max Score", "Notes"]):
            fill_cell(cell, label, bold=True, color=SN_WHITE, size=10, shading="293E40")

        for i, w in enumerate(weights):
            row = table.add_row()
            is_total = w.get("tier", "").lower() == "total"
            bg = "E8E8E8" if is_total else ("F5F5F5" if i % 2 == 0 else "FFFFFF")
            fill_cell(row.cells[0], w.get("tier", ""), bold=is_total, size=10, shading=bg)
            fill_cell(row.cells[1], w.get("max", ""), bold=is_total, size=10,
                      align=WD_ALIGN_PARAGRAPH.CENTER, shading=bg)
            fill_cell(row.cells[2], w.get("note", ""), size=9, shading=bg)

        set_col_widths(table, [3.0, 1.2, 2.8])

        # Hard rule note
        para(doc,
             "Zero on any Trust & Governance criterion warrants disqualification "
             "from further evaluation.",
             size=9, italic=True, color=SN_GRAY, before=10, after=0,
             indent=0)


# ── Validation ─────────────────────────────────────────────────────────────
def validate_spec(spec):
    errors = []
    if "slides" not in spec:
        errors.append("Missing 'slides' key")
    for i, slide in enumerate(spec.get("slides", [])):
        if "type" not in slide:
            errors.append(f"Slide {i}: missing 'type'")
    if errors:
        for e in errors:
            print(f"[SPEC ERROR] {e}", file=sys.stderr)
        sys.exit(1)


def validate_output(path):
    try:
        d = Document(str(path))
        count = len(d.paragraphs) + sum(len(t.rows) for t in d.tables)
        print(f"[OK] {path}  ({count} elements)")
    except Exception as exc:
        print(f"[ERROR] Output invalid: {exc}", file=sys.stderr)
        sys.exit(1)


def check_repair(path):
    p = Path(path)
    if not p.exists():
        print(f"[ERROR] File not found: {p}", file=sys.stderr)
        sys.exit(1)
    try:
        d = Document(str(p))
        print(f"[OK] {p.name}  paragraphs={len(d.paragraphs)}  tables={len(d.tables)}")
    except Exception as exc:
        print(f"[CORRUPT] {p.name}: {exc}", file=sys.stderr)
        sys.exit(1)


# ── Main ───────────────────────────────────────────────────────────────────
RENDERERS = {
    "cover_green":       render_cover,
    "four_section_grid": render_dimensions_overview,
    "scoring_table":     render_scoring_table,
    "scoring_guide":     render_scoring_guide,
}


def main():
    args = sys.argv[1:]

    if not args:
        print(__doc__)
        sys.exit(0)

    if args[0] == "--check-repair":
        if len(args) < 2:
            print("Usage: render_doc.py --check-repair <file.docx>", file=sys.stderr)
            sys.exit(1)
        check_repair(args[1])
        return

    spec_path = Path(args[0])
    output_override = None
    if "--output" in args:
        idx = args.index("--output")
        output_override = Path(args[idx + 1])

    with open(spec_path) as f:
        spec = yaml.safe_load(f)

    validate_spec(spec)

    if output_override:
        out_path = output_override
    elif "doc_output" in spec:
        out_path = Path(spec["doc_output"])
    else:
        deck_out = spec.get("output", "")
        out_path = (Path(deck_out).with_suffix(".docx")
                    if deck_out else spec_path.with_suffix(".docx"))

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.1)
        section.right_margin = Inches(1.1)

    for slide in spec.get("slides", []):
        renderer = RENDERERS.get(slide.get("type"))
        if renderer:
            renderer(doc, slide)
        else:
            print(f"[WARN] Unknown slide type: {slide.get('type')}")

    doc.save(str(out_path))
    print(f"[DONE] Saved: {out_path}")
    validate_output(out_path)


if __name__ == "__main__":
    main()
